"""
DESIGN Exports
Server-side Word, PDF and Excel generation for DESIGN outputs. Word and PDF render from a small
document model: a list of ("h1"|"h2"|"p"|"label"|"bullets"|"meta"|"table", ...) elements.
"""

import io
import json
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Pt, RGBColor, Inches
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MIME = "application/pdf"

LABEL_TEXT = {
    "fact": "Source fact",
    "verified": "Verified source fact",
    "pending": "Source fact – pending verification",
    "calculated": "Calculated result",
    "ai_interpretation": "AI interpretation",
    "unresolved": "Unresolved issue",
    "note": "Note",
}


# ─── Word ────────────────────────────────────────────────────────────

def to_docx(model: list, title: str, wide: bool = False) -> bytes:
    doc = Document()
    if wide:
        sec = doc.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Inches(0.5))
    doc.add_heading(title, level=0)
    for el in model:
        kind = el[0]
        if kind == "h1":
            doc.add_heading(el[1], level=1)
        elif kind == "h2":
            doc.add_heading(el[1], level=2)
        elif kind == "p":
            doc.add_paragraph(el[1])
        elif kind == "label":
            _, label, text, cites = el
            p = doc.add_paragraph()
            r = p.add_run(f"[{LABEL_TEXT.get(label, label)}] ")
            r.bold = True
            r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            p.add_run(text)
            if cites:
                c = p.add_run(f"  ({'; '.join(cites)})")
                c.italic = True
                c.font.size = Pt(8)
        elif kind == "bullets":
            for item in el[1]:
                doc.add_paragraph(str(item), style="List Bullet")
        elif kind == "meta":
            t = doc.add_table(rows=0, cols=2)
            t.style = "Light Grid Accent 1"
            for k, v in el[1]:
                row = t.add_row().cells
                row[0].text, row[1].text = str(k), str(v)
        elif kind == "table":
            _, columns, rows = el[:3]
            t = doc.add_table(rows=1, cols=len(columns))
            t.style = "Table Grid"
            for i, c in enumerate(columns):
                cell = t.rows[0].cells[i]
                cell.text = str(c)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                cells = t.add_row().cells
                for i, v in enumerate(row[:len(columns)]):
                    cells[i].text = "" if v is None else str(v)
            if wide:
                for row in t.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            for r in p.runs:
                                r.font.size = Pt(7)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─── PDF ─────────────────────────────────────────────────────────────

def to_pdf(model: list, title: str, wide: bool = False) -> bytes:
    buf = io.BytesIO()
    size = landscape(A4) if wide else A4
    doc = SimpleDocTemplate(buf, pagesize=size, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36, title=title)
    ss = getSampleStyleSheet()
    st_title = ParagraphStyle("t", parent=ss["Heading1"], fontSize=18, leading=22, textColor=colors.HexColor("#0f172a"))
    st_h1 = ParagraphStyle("h1", parent=ss["Heading2"], fontSize=13, textColor=colors.HexColor("#1e3a8a"), spaceBefore=10)
    st_h2 = ParagraphStyle("h2", parent=ss["Heading3"], fontSize=11, textColor=colors.HexColor("#2563eb"))
    st_body = ParagraphStyle("b", parent=ss["Normal"], fontSize=9, leading=12)
    st_cell = ParagraphStyle("c", parent=ss["Normal"], fontSize=7 if wide else 8, leading=9 if wide else 10)
    story = [Paragraph(escape(title), st_title), Spacer(1, 6)]
    for el in model:
        kind = el[0]
        if kind == "h1":
            story.append(Paragraph(escape(el[1]), st_h1))
        elif kind == "h2":
            story.append(Paragraph(escape(el[1]), st_h2))
        elif kind == "p":
            story.append(Paragraph(escape(el[1]), st_body))
        elif kind == "label":
            _, label, text, cites = el
            cite = f" <i><font size=7>({escape('; '.join(cites))})</font></i>" if cites else ""
            story.append(Paragraph(f"<b><font color='#1e3a8a'>[{escape(LABEL_TEXT.get(label, label))}]</font></b> {escape(text)}{cite}", st_body))
        elif kind == "bullets":
            for item in el[1]:
                story.append(Paragraph("&bull; " + escape(str(item)), st_body))
        elif kind in ("meta", "table"):
            if kind == "meta":
                data = [[Paragraph(f"<b>{escape(str(k))}</b>", st_cell), Paragraph(escape(str(v)), st_cell)] for k, v in el[1]]
                widths = [140, size[0] - 72 - 140]
            else:
                columns, rows = el[1], el[2]
                data = [[Paragraph(f"<b>{escape(str(c))}</b>", st_cell) for c in columns]]
                data += [[Paragraph(escape("" if v is None else str(v)), st_cell) for v in r[:len(columns)]] for r in rows]
                avail = size[0] - 72
                first = min(160, avail * 0.3) if len(columns) > 4 else avail / len(columns)
                rest = (avail - first) / max(1, len(columns) - 1)
                widths = [first] + [rest] * (len(columns) - 1)
            t = Table(data, colWidths=widths, repeatRows=1 if kind == "table" else 0)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BACKGROUND", (0, 0), (-1, 0) if kind == "table" else (0, -1), colors.HexColor("#eff6ff")),
            ]))
            story += [t, Spacer(1, 6)]
        story.append(Spacer(1, 2))
    doc.build(story)
    return buf.getvalue()


def render(model: list, title: str, fmt: str, wide: bool = False) -> tuple[bytes, str, str]:
    if fmt == "docx":
        return to_docx(model, title, wide), DOCX_MIME, "docx"
    if fmt == "pdf":
        return to_pdf(model, title, wide), PDF_MIME, "pdf"
    raise ValueError("Unsupported format")


# ─── Excel ───────────────────────────────────────────────────────────

_HEAD = Font(bold=True, color="FFFFFF")
_FILL = PatternFill("solid", fgColor="1E3A8A")


def _sheet(ws, columns: list[str], rows: list[list], widths: list[int] | None = None, start_row: int = 1):
    for i, c in enumerate(columns, start=1):
        cell = ws.cell(row=start_row, column=i, value=c)
        cell.font, cell.fill = _HEAD, _FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r_i, r in enumerate(rows, start=start_row + 1):
        for c_i, v in enumerate(r, start=1):
            cell = ws.cell(row=r_i, column=c_i, value=v if not isinstance(v, (list, dict)) else json.dumps(v))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _meta_sheet(wb, meta: list[tuple[str, str]]):
    ws = wb.active
    ws.title = "About"
    _sheet(ws, ["Item", "Value"], [list(m) for m in meta], [28, 80])


def _val(v) -> str:
    return "\n".join(str(x) for x in v) if isinstance(v, list) else ("" if v is None else str(v))


def _ev(evs: list[dict]) -> str:
    return "\n".join(
        ", ".join(x for x in [f"p.{e['page']}" if e.get("page") else "", e.get("section") or "", e.get("block_id") or ""] if x)
        + (f": \"{e.get('quote')}\"" if e.get("quote") else "") + ("" if e.get("verified", True) else " [NOT FOUND IN SOURCE]")
        for e in evs or [])


def pi_to_xlsx(meta: list[tuple[str, str]], data: dict) -> bytes:
    wb = Workbook()
    _meta_sheet(wb, meta)
    rows, review = [], []
    for f in data["fields"]:
        v = f["value"] if f["status"] not in ("missing", "not_requested") else (
            "Not specified" if f["status"] == "missing" else "Not extracted (excluded by preferences)")
        row = [f["category"], f["label"], _val(v), f["status"], f["confidence"], f["confidence_level"],
               f["verification"]["status"], f.get("source") or "", "; ".join(f.get("flags", [])), _ev(f.get("evidence"))]
        rows.append(row)
        if f["status"] not in ("extracted", "not_requested") or (f["confidence_level"] == "low" and f["status"] != "not_requested"):
            review.append(row)
    cols = ["Category", "Field", "Value", "Extraction status", "Confidence", "Level", "Verification", "Source", "Flags", "Evidence"]
    widths = [18, 26, 60, 16, 11, 9, 13, 9, 40, 70]
    if data.get("characteristics"):
        _sheet(wb.create_sheet("Study characteristics"), ["Characteristic", "Value", "Extraction status", "Verification"],
               [[c["label"], c["value"], c["status"], c["verification"]] for c in data["characteristics"]], [26, 90, 16, 13])
    _sheet(wb.create_sheet("Fields"), cols, rows, widths)
    _sheet(wb.create_sheet("Missing & low confidence"), cols, review, widths)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _timing_label(v: dict) -> str:
    t = v.get("timing")
    return t["label"] if t else "Not specified"


def _window_label(v: dict) -> str:
    w = v.get("window")
    if not w:
        return "Not specified"
    return f"-{w['minus']}/+{w['plus']} {w['unit']}{'s' if w['plus'] != 1 else ''}"


def soa_rows(matrix: dict) -> tuple[list[str], list[list[str]]]:
    columns = ["Assessment"] + [v["label"] for v in matrix["visits"]]
    rows = [["Timing"] + [_timing_label(v) for v in matrix["visits"]],
            ["Visit window"] + [_window_label(v) for v in matrix["visits"]]]
    category = object()
    for a in matrix["assessments"]:
        if a.get("category") and a.get("category") != category:
            category = a["category"]
            rows.append([f"— {category} —"] + [""] * len(matrix["visits"]))
        row = [a["name"]]
        for v in matrix["visits"]:
            c = matrix["cells"].get(f"{a['id']}|{v['id']}")
            if not c:
                row.append("")
            else:
                refs = ",".join(c.get("footnotes") or [])
                row.append(c["mark"] + (f" ({refs})" if refs else ""))
        rows.append(row)
    return columns, rows


def soa_to_xlsx(meta: list[tuple[str, str]], matrix: dict) -> bytes:
    wb = Workbook()
    _meta_sheet(wb, meta)
    ws = wb.create_sheet("Schedule of Assessments")
    columns, rows = soa_rows(matrix)
    _sheet(ws, columns, rows, [36] + [14] * len(matrix["visits"]))
    r = len(rows) + 3
    ws.cell(row=r, column=1, value="Footnotes").font = Font(bold=True)
    for f in matrix.get("footnotes", []):
        r += 1
        ws.cell(row=r, column=1, value=f["marker"])
        ws.cell(row=r, column=2, value=f["text"])
    r += 2
    ws.cell(row=r, column=1, value="Timing conventions").font = Font(bold=True)
    for c in matrix.get("conventions", []):
        r += 1
        ws.cell(row=r, column=1, value=c["text"])
    ws.freeze_panes = "B4"
    _sheet(wb.create_sheet("Visits"), ["ID", "Label", "Visit no.", "Phase", "Timing (as stated)", "Window", "Source"],
           [[v["id"], v["label"], v.get("visit_number") or "", v.get("phase"), _timing_label(v), _window_label(v),
             _ev([v["source"]]) if v.get("source") else ""] for v in matrix["visits"]], [6, 30, 9, 16, 20, 18, 40])
    _sheet(wb.create_sheet("Timeline"), ["Order", "Visit", "Phase", "Timing (as stated)", "Approx. study day", "Window"],
           [[n, t["label"], t["phase"], t["timing"], t["approx_day"] if t["approx_day"] is not None else "Not specified", t["window"]]
            for n, t in enumerate(timeline(matrix), start=1)], [7, 34, 16, 20, 16, 18])
    _sheet(wb.create_sheet("Issues"), ["ID", "Type", "Severity", "Status", "Description", "Resolution"],
           [[i["id"], i["type"], i["severity"], i["status"], i["description"],
             (i.get("resolution") or {}).get("note", "")] for i in matrix.get("issues", [])], [6, 26, 10, 10, 70, 40])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_DAYS_PER = {"day": 1, "week": 7, "month": 30.4375, "year": 365.25}


def timeline(matrix: dict) -> list[dict]:
    """Visits in chronological order. The approximate study day is derived from the stated timing only for ordering."""
    rows = []
    for idx, v in enumerate(matrix["visits"]):
        t = v.get("timing")
        day = round(t["value"] * _DAYS_PER[t["unit"]]) if t and t["unit"] in _DAYS_PER else None
        rows.append({"visit_id": v["id"], "label": v["label"], "phase": v.get("phase") or "Not specified",
                     "timing": _timing_label(v), "approx_day": day, "approx": bool(t and t["unit"] != "day"),
                     "window": _window_label(v), "order": idx})
    return sorted(rows, key=lambda r: (r["approx_day"] is None, r["approx_day"] if r["approx_day"] is not None else 0, r["order"]))


# ─── Protocol-specific SoA templates ─────────────────────────────────

TEMPLATE_PLACEHOLDERS = ["{{SOA_TABLE}}", "{{SOA_FOOTNOTES}}", "{{SOA_CONVENTIONS}}", "{{STUDY_ID}}", "{{STUDY_TITLE}}",
                         "{{PROTOCOL_VERSION}}", "{{STATUS}}", "{{REVISION}}", "{{EXPORTED}}"]


def template_placeholders(content: bytes, fmt: str) -> list[str]:
    """Placeholders present in an uploaded template (used to validate it on upload)."""
    text = ""
    if fmt == "xlsx":
        wb = load_workbook(io.BytesIO(content))
        text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)
    else:
        doc = Document(io.BytesIO(content))
        text = " ".join(p.text for p in doc.paragraphs) + " " + " ".join(
            c.text for t in doc.tables for r in t.rows for c in r.cells)
    return [p for p in TEMPLATE_PLACEHOLDERS if p in text]


def _replace_values(values: dict[str, str], text: str) -> str:
    for k, v in values.items():
        text = text.replace(k, v)
    return text


def soa_fill_xlsx_template(content: bytes, values: dict[str, str], matrix: dict) -> bytes:
    """Writes the matrix at {{SOA_TABLE}} keeping the template's styles; other placeholders are replaced."""
    from copy import copy
    wb = load_workbook(io.BytesIO(content))
    columns, rows = soa_rows(matrix)
    anchor = None
    lists = {"{{SOA_FOOTNOTES}}": [f"{f['marker']}: {f['text']}" for f in matrix.get("footnotes", [])],
             "{{SOA_CONVENTIONS}}": [c["text"] for c in matrix.get("conventions", [])]}
    list_anchors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value, str):
                    continue
                if "{{SOA_TABLE}}" in c.value:
                    anchor = (ws, c)
                for key in lists:
                    if key in c.value:
                        list_anchors.append((ws, c, key))
                c.value = _replace_values(values, c.value)
    if anchor is None:
        ws = wb.create_sheet("Schedule of Assessments")
        _sheet(ws, columns, rows, [36] + [14] * len(matrix["visits"]))
    else:
        ws, cell = anchor
        style = copy(cell._style)
        for r_off, line in enumerate([columns] + rows):
            for c_off, val in enumerate(line):
                target = ws.cell(row=cell.row + r_off, column=cell.column + c_off, value=val)
                target._style = copy(style)
                if r_off == 0:
                    target.font = Font(bold=True)
    for ws, cell, key in list_anchors:
        for n, line in enumerate(lists[key] or ["None"]):
            ws.cell(row=cell.row + n, column=cell.column, value=line)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def soa_fill_docx_template(content: bytes, values: dict[str, str], matrix: dict) -> bytes:
    """Inserts the matrix table at the {{SOA_TABLE}} paragraph; other placeholders are replaced in place."""
    doc = Document(io.BytesIO(content))
    columns, rows = soa_rows(matrix)
    lists = {"{{SOA_FOOTNOTES}}": [f"{f['marker']}: {f['text']}" for f in matrix.get("footnotes", [])],
             "{{SOA_CONVENTIONS}}": [c["text"] for c in matrix.get("conventions", [])]}

    def fill_paragraph(p):
        if "{{SOA_TABLE}}" in p.text:
            table = doc.add_table(rows=1, cols=len(columns))
            table.style = "Table Grid"
            for i, col in enumerate(columns):
                table.rows[0].cells[i].text = str(col)
            for r in rows:
                cells = table.add_row().cells
                for i, v in enumerate(r[:len(columns)]):
                    cells[i].text = "" if v is None else str(v)
            p._p.addnext(table._tbl)
            p.text = p.text.replace("{{SOA_TABLE}}", "")
            return
        for key, items in lists.items():
            if key in p.text:
                p.text = p.text.replace(key, "\n".join(items) or "None")
        if any(k in p.text for k in values):
            p.text = _replace_values(values, p.text)

    for p in list(doc.paragraphs):
        fill_paragraph(p)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    if "{{SOA_TABLE}}" not in p.text:
                        fill_paragraph(p)
    for section in doc.sections:
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                if "{{SOA_TABLE}}" not in p.text:
                    fill_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def soa_model(meta: list[tuple[str, str]], matrix: dict) -> list:
    columns, rows = soa_rows(matrix)
    model = [("meta", meta), ("h1", "Schedule of Assessments"), ("table", columns, rows),
             ("h2", "Visit timeline"),
             ("table", ["#", "Visit", "Phase", "Timing (as stated)", "Approx. study day", "Window"],
              [[n, t["label"], t["phase"], t["timing"], t["approx_day"] if t["approx_day"] is not None else "Not specified", t["window"]]
               for n, t in enumerate(timeline(matrix), start=1)])]
    if matrix.get("footnotes"):
        model += [("h2", "Footnotes"), ("bullets", [f"{f['marker']}: {f['text']}" for f in matrix["footnotes"]])]
    if matrix.get("conventions"):
        model += [("h2", "Timing conventions (as stated in the protocol)"), ("bullets", [c["text"] for c in matrix["conventions"]])]
    open_issues = [i for i in matrix.get("issues", []) if i["status"] != "resolved"]
    if open_issues:
        model += [("h2", "Unresolved scheduling issues"), ("bullets", [f"[{i['severity']}] {i['description']}" for i in open_issues])]
    return model


def sample_size_model(meta: list[tuple[str, str]], result: dict) -> list:
    arms = result["arms"]
    model = [("meta", meta), ("h1", "Result")]
    rows = [[arm, result["evaluable"]["per_arm"][i], result["enrollment"]["per_arm"][i]] for i, arm in enumerate(arms)]
    rows.append(["Total", result["evaluable"]["total"], result["enrollment"]["total"]])
    model.append(("table", ["Group", "Evaluable", "Enrollment target"], rows))
    if result.get("events"):
        model.append(("label", "calculated", f"Required number of events: {result['events']}", []))
    model += [
        ("h1", "Method"),
        ("p", f"{result['method_name']} — {result['reference']}"),
        ("p", f"Statistical test: {result.get('statistical_test', '')}. Default justification: {result.get('test_justification', '')}"),
        ("p", f"Formula: {result['formula']}"),
        ("p", f"Attrition adjustment: {result['attrition_formula']}"),
        ("p", f"Rounding: {result['rounding']}"),
        ("p", f"Unrounded values: {json.dumps(result['raw'])}"),
        ("h1", "Assumptions"), ("bullets", result["assumptions"]),
        ("h1", "Sensitivity analysis"),
        ("table", ["Parameter varied", "Value", "Evaluable total", "Enrollment total", "Events"],
         [[r["parameter"], r["value"], r.get("evaluable_total", r.get("error")), r.get("enrollment_total", ""), r.get("events") or ""]
          for r in result["sensitivity"]]),
        ("h1", "Limitations"), ("bullets", result["limitations"] + result.get("warnings", [])),
        ("h1", "Inputs (versioned)"), ("p", json.dumps(result["inputs"], indent=1)),
        ("label", "note", result["disclaimer"], []),
    ]
    return model


def _refs(evs: list[dict]) -> str:
    return "; ".join(", ".join(x for x in [f"p.{e['page']}" if e.get("page") else "", e.get("section") or "", e.get("block_id") or ""] if x)
                     for e in evs or []) or "—"


def review_to_xlsx(meta: list[tuple[str, str]], data: dict) -> bytes:
    wb = Workbook()
    _meta_sheet(wb, meta)
    _sheet(wb.create_sheet("Findings"),
           ["ID", "Severity", "Priority", "Category", "Title", "Finding", "Recommendation", "Source", "Origin", "Certainty", "Status", "Resolution"],
           [[f["id"], f["severity"], f["priority"], f["category_label"], f["title"], f["description"], f["recommendation"],
             _refs(f["evidence"]) + ("\n" + "\n".join(f'"{e.get("quote")}"' for e in f["evidence"][:2]) if f["evidence"] else ""),
             "AI" if f["origin"] == "ai" else "Rule", f["certainty"], f["status"], (f.get("resolution") or {}).get("note", "")]
            for f in data["findings"]], [6, 10, 9, 18, 36, 60, 50, 40, 7, 10, 11, 40])
    _sheet(wb.create_sheet("Completeness checklist"), ["Element", "Checked in", "Status", "Source"],
           [[c["element"], c["source"], c["status"].replace("_", " "), _refs(c.get("evidence"))] for c in data["checklist"]], [40, 22, 16, 50])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def review_model(meta: list[tuple[str, str]], data: dict) -> list:
    s = data["summary"]
    comp = s["completeness"]
    model = [("meta", meta), ("h1", "Summary"),
             ("p", f"{s['total']} findings ({s['open']} open): " + ", ".join(f"{v} {k}" for k, v in s["by_severity"].items())
              + ". By category: " + ", ".join(f"{v} {k}" for k, v in s["by_category"].items() if v) + "."),
             ("p", f"Completeness: {comp['present']} of {comp['checked']} expected protocol elements present"
              + (f" ({comp['percent']}%)." if comp["percent"] is not None else "."))]
    if data["basis"].get("warning"):
        model.append(("label", "note", data["basis"]["warning"], []))
    model.append(("h1", "Findings"))
    for f in data["findings"]:
        model.append(("h2", f"{f['id']} · {f['severity'].upper()} · {f['category_label']} — {f['title']}"))
        model.append(("label", "ai_interpretation" if f["origin"] == "ai" else "unresolved" if f["status"] == "open" else "note",
                      f"{f['description']} Recommendation: {f['recommendation']} [Status: {f['status']}"
                      + (f" — {f['resolution']['note']}" if f.get("resolution") else "") + "]",
                      [_refs([e]) + (f' "{e.get("quote")}"' if e.get("quote") else "") for e in f["evidence"][:3]]))
    model += [("h1", "Completeness checklist"),
              ("table", ["Element", "Checked in", "Status", "Source"],
               [[c["element"], c["source"], c["status"].replace("_", " "), _refs(c.get("evidence"))] for c in data["checklist"]])]
    return model
