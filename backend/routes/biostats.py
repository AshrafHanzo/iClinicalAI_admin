"""
Biostats Routes
FastAPI endpoints for iClinicalAI ANALYSE module features.
"""

import io
from datetime import datetime
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from database import get_db, UserDB
from routes.auth import get_current_user
import services.db_service as db_service
from services.analyse_engine import run_analyse_step

# Word Export
from docx import Document

# PDF Export
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

router = APIRouter(prefix="/api/biostats", tags=["Biostatistics"])


class BiostatsAnalysisRequest(BaseModel):
    step: str
    doc_id: str | None = None
    extra_input: str | None = None


class ExportRequest(BaseModel):
    title: str
    content: str
    filename: str


@router.post("/analyze")
async def analyze_biostats_step(
    request: BiostatsAnalysisRequest, 
    db: Session = Depends(get_db),
    current_user: UserDB = Depends(get_current_user)
):
    """Run an AI-powered biostatistics analysis step."""
    document_text = ""
    doc = None
    if request.doc_id:
        doc = db_service.get_document(db, request.doc_id, user_id=current_user.id)
        if doc:
            document_text = doc["extracted_text"]
        else:
            raise HTTPException(status_code=404, detail="Document not found")
    else:
        document_text = "No study protocol selected. Operating under standard clinical trial biostatistics guidelines."

    try:
        result = await run_analyse_step(
            step=request.step,
            document_text=document_text,
            extra_input=request.extra_input
        )
        if doc:
            doc.setdefault("analysis_results", {})[request.step] = result
            db_service.save_document(db, doc, user_id=current_user.id)
            
        db_service.log_api_usage(db, current_user.id, "ANALYSE", 13400)
        return {
            "step": request.step,
            "result": result,
            "timestamp": datetime.now().isoformat()
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Biostatistics analysis failed: {str(e)}")


@router.post("/export/word")
async def export_to_word(request: ExportRequest):
    """Export markdown-like content to a Word Document."""
    doc = Document()
    doc.add_heading(request.title, level=0)

    lines = request.content.split("\n")
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:], style='List Bullet')
        else:
            doc.add_paragraph(line)

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={request.filename}.docx"}
    )


@router.post("/export/pdf")
async def export_to_pdf(request: ExportRequest):
    """Export content to a PDF Document using ReportLab."""
    pdf_buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        pdf_buffer,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()
    
    # Custom Styles for Premium Look
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=24,
        leading=28,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=20
    )
    
    h1_style = ParagraphStyle(
        'H1',
        parent=styles['Heading2'],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#1e3a8a'),
        spaceBefore=14,
        spaceAfter=6
    )
    
    h2_style = ParagraphStyle(
        'H2',
        parent=styles['Heading3'],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#2563eb'),
        spaceBefore=10,
        spaceAfter=4
    )
    
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155'),
        spaceAfter=8
    )

    story = []
    story.append(Paragraph(request.title, title_style))
    story.append(Spacer(1, 12))

    lines = request.content.split("\n")
    for line in lines:
        line = line.strip()
        if not line:
            continue
        
        # Simple markdown headers parsing
        if line.startswith("# "):
            story.append(Paragraph(line[2:], h1_style))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], h2_style))
        elif line.startswith("### "):
            story.append(Paragraph(line[4:], h2_style))
        elif line.startswith("- ") or line.startswith("* "):
            bullet_text = f"&bull; {line[2:]}"
            story.append(Paragraph(bullet_text, body_style))
        else:
            story.append(Paragraph(line, body_style))

    doc.build(story)
    pdf_buffer.seek(0)

    return StreamingResponse(
        pdf_buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={request.filename}.pdf"}
    )
