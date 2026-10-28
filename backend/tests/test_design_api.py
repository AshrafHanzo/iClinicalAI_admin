"""End-to-end API tests for the DESIGN workspace: Frontend contract → API → services → DB → response."""

import hashlib
import io
import json

import pytest
from docx import Document
from openpyxl import load_workbook

import services.design_reports as reports_service
from tests.conftest import make_user
from tests.fixtures import build_docx, build_pdf

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def create_study(client, headers, code="DX-301"):
    r = client.post("/api/design/studies", json={"study_code": code, "title": "Drug X in T2D", "sponsor": "Acme"}, headers=headers)
    assert r.status_code == 200, r.text
    return r.json()


def upload(client, headers, study_id, content, name="protocol.docx", **form):
    files = {"file": (name, content, DOCX if name.endswith("docx") else "application/pdf")}
    data = {"preferences": json.dumps({"use_ai": False}), **form}
    return client.post(f"/api/design/studies/{study_id}/versions", files=files, data=data, headers=headers)


@pytest.fixture
def study_v1(client, owner):
    _, h = owner
    s = create_study(client, h)
    r = upload(client, h, s["id"], build_docx(), version_label="1.0")
    assert r.status_code == 200, r.text
    v = client.get(f"/api/design/versions/{r.json()['id']}", headers=h).json()
    assert v["processing_status"] == "completed", v["processing_message"]
    return s, v, h


def add_member(client, owner_headers, study_id, role):
    email, h = make_user()
    r = client.post(f"/api/design/studies/{study_id}/members", json={"email": email, "role": role}, headers=owner_headers)
    assert r.status_code == 200, r.text
    return h


# ─── Auth & isolation ────────────────────────────────────────────────

def test_requires_authentication(client):
    assert client.get("/api/design/studies").status_code == 401


def test_study_isolation(client, study_v1, other):
    s, v, _ = study_v1
    _, oh = other
    assert client.get(f"/api/design/studies/{s['id']}", headers=oh).status_code == 404
    assert client.get(f"/api/design/versions/{v['id']}/intelligence", headers=oh).status_code == 404
    assert client.post(f"/api/design/versions/{v['id']}/assistant", json={"question": "phase?"}, headers=oh).status_code == 404
    assert client.get(f"/api/design/versions/{v['id']}/original", headers=oh).status_code == 404
    assert s["id"] not in [x["id"] for x in client.get("/api/design/studies", headers=oh).json()["studies"]]


def test_viewer_can_read_but_not_edit(client, study_v1):
    s, v, h = study_v1
    vh = add_member(client, h, s["id"], "viewer")
    assert client.get(f"/api/design/versions/{v['id']}/intelligence", headers=vh).status_code == 200
    r = client.post(f"/api/design/versions/{v['id']}/intelligence/fields", json={"updates": [{"key": "phase", "value": "Phase 2"}]}, headers=vh)
    assert r.status_code == 403
    assert upload(client, vh, s["id"], build_docx("v2")).status_code == 403


# ─── Feature 01: Protocol Intelligence ───────────────────────────────

def test_upload_rejects_invalid_files_without_creating_versions(client, owner):
    _, h = owner
    s = create_study(client, h, "BAD-1")
    r = upload(client, h, s["id"], b"plain text", name="notes.txt")
    assert r.status_code == 400 and "PDF or DOCX" in r.json()["detail"]
    r = upload(client, h, s["id"], b"%PDF-1.4 broken", name="broken.pdf")
    assert r.status_code == 400
    assert client.get(f"/api/design/studies/{s['id']}", headers=h).json()["versions"] == []


def test_unreadable_pdf_shows_failed_status(client, owner):
    _, h = owner
    s = create_study(client, h, "SCAN-1")
    r = upload(client, h, s["id"], build_pdf(blank=True), name="scan.pdf")
    assert r.status_code == 200
    v = client.get(f"/api/design/versions/{r.json()['id']}", headers=h).json()
    assert v["processing_status"] == "failed" and "No readable text" in v["processing_message"]


def test_extraction_is_persisted_with_evidence(client, study_v1):
    s, v, h = study_v1
    r = client.get(f"/api/design/versions/{v['id']}/intelligence", headers=h)
    assert r.status_code == 200
    a = r.json()
    fields = {f["key"]: f for f in a["data"]["fields"]}
    assert fields["phase"]["value"] == "Phase 3" and fields["phase"]["evidence"]
    assert fields["hypotheses"]["value"] == "Not specified"
    assert a["protocol_version_id"] == v["id"] and a["study_id"] == s["id"]
    # Evidence viewer returns the cited block with context
    bid = fields["sample_size"]["evidence"][0]["block_id"]
    blk = client.get(f"/api/design/versions/{v['id']}/blocks/{bid}", headers=h).json()
    assert "240 participants" in blk["block"]["text"] and blk["before"]


def test_original_document_is_unchanged(client, study_v1):
    _, v, h = study_v1
    r = client.get(f"/api/design/versions/{v['id']}/original", headers=h)
    assert r.status_code == 200
    assert hashlib.sha256(r.content).hexdigest() == v["file_sha256"]


def test_edit_and_verify_workflow(client, study_v1):
    s, v, h = study_v1
    eh = add_member(client, h, s["id"], "editor")
    url = f"/api/design/versions/{v['id']}/intelligence/fields"
    r = client.post(url, json={"updates": [{"key": "study_duration", "value": "28 weeks", "note": "From section 3"}]}, headers=eh)
    assert r.status_code == 200
    f = {x["key"]: x for x in r.json()["data"]["fields"]}["study_duration"]
    assert f["value"] == "28 weeks" and f["source"] == "user" and f["verification"]["status"] == "corrected"
    assert f["history"][0]["value"] == "Not specified"
    # Editors cannot verify; reviewers/owners can
    assert client.post(url, json={"updates": [{"key": "study_duration", "verify": True}]}, headers=eh).status_code == 403
    r = client.post(url, json={"updates": [{"key": "study_duration", "verify": True}, {"key": "phase", "verify": True}]}, headers=h)
    assert r.status_code == 200
    body = r.json()
    assert body["revision"] == 3 and len(body["history"]) == 3
    fields = {x["key"]: x for x in body["data"]["fields"]}
    assert fields["phase"]["verification"]["status"] == "verified"
    assert client.get(f"/api/design/versions/{v['id']}", headers=h).json()["review_status"] == "in_review"
    # Stale writes are rejected instead of overwriting
    r = client.post(url, json={"updates": [{"key": "phase", "value": "Phase 2"}], "base_revision": 1}, headers=h)
    assert r.status_code == 409
    audit = client.get(f"/api/design/studies/{s['id']}/audit", headers=h).json()["entries"]
    assert any(e["action"] == "pi_fields_updated" for e in audit)


def test_intelligence_exports(client, study_v1):
    s, v, h = study_v1
    r = client.get(f"/api/design/versions/{v['id']}/intelligence/export?format=json", headers=h)
    body = json.loads(r.content)
    assert body["metadata"]["Study ID"] == "DX-301" and body["metadata"]["Protocol version"] == "1.0"
    r = client.get(f"/api/design/versions/{v['id']}/intelligence/export?format=xlsx", headers=h)
    wb = load_workbook(io.BytesIO(r.content))
    assert {"About", "Fields", "Missing & low confidence"} <= set(wb.sheetnames)
    values = [c.value for c in wb["Missing & low confidence"]["B"]]
    assert "Study hypotheses" in values


def test_registered_document_cannot_be_deleted(client, study_v1):
    _, v, h = study_v1
    r = client.delete(f"/api/documents/{v['document_id']}", headers=h)
    assert r.status_code == 409


# ─── Feature 06: Endpoint Intelligence ───────────────────────────────

def test_endpoint_intelligence(client, study_v1):
    s, v, h = study_v1
    r = client.post(f"/api/design/versions/{v['id']}/endpoints/analyze", json={"use_ai": False}, headers=h)
    assert r.status_code == 200, r.text
    d = r.json()["data"]
    ep1 = d["register"][0]
    assert ep1["type"] == "primary" and ep1["evidence"][0]["section"] == "2.3 Primary Endpoint"
    assert {"endpoint_id": "EP1", "objective_id": "OBJ1"}.items() <= next(l for l in d["mapping"] if l["endpoint_id"] == "EP1").items()
    types = {(i["type"], i["certainty"]) for i in d["issues"]}
    assert ("inconsistent_timepoint", "confirmed") in types  # Week 24 vs Week 26
    assert ("missing_timepoint", "confirmed") in types
    assert ("missing_measurement", "confirmed") in types
    assert ("objective_without_endpoint", "potential") in types
    assert d["basis"]["warning"]  # analysis based on unverified fields is disclosed
    hierarchy = d["hierarchy"]

    sug = next(x for x in d["suggestions"] if x["requires_scientific_review"])
    eh = add_member(client, h, s["id"], "editor")
    url = f"/api/design/versions/{v['id']}/endpoints/suggestions/{sug['id']}/review"
    assert client.post(url, json={"action": "comment", "comment": "Check SAP"}, headers=eh).status_code == 200
    assert client.post(url, json={"action": "approve"}, headers=eh).status_code == 403
    r = client.post(url, json={"action": "approve", "comment": "Agreed"}, headers=h)
    d2 = r.json()["data"]
    s2 = next(x for x in d2["suggestions"] if x["id"] == sug["id"])
    assert s2["status"] == "approved" and s2["comments"][0]["text"] == "Check SAP"
    assert d2["hierarchy"] == hierarchy and d2["register"] == d["register"]  # never changed automatically

    # Re-running keeps review decisions
    d3 = client.post(f"/api/design/versions/{v['id']}/endpoints/analyze", json={"use_ai": False}, headers=h).json()
    assert next(x for x in d3["data"]["suggestions"] if x["id"] == sug["id"])["status"] == "approved"
    assert d3["revision"] == 4


def test_endpoint_requires_protocol_intelligence(client, owner):
    _, h = owner
    s = create_study(client, h, "EMPTY-1")
    r = upload(client, h, s["id"], build_pdf(blank=True), name="scan.pdf")
    assert client.post(f"/api/design/versions/{r.json()['id']}/endpoints/analyze", json={"use_ai": False}, headers=h).status_code == 409


# ─── Feature 07: Schedule of Assessments ─────────────────────────────

def test_schedule_of_assessments(client, study_v1):
    s, v, h = study_v1
    r = client.post(f"/api/design/versions/{v['id']}/soa/generate", json={"use_ai": False}, headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    m = body["data"]
    assert [x["label"].split(" / ")[0] for x in m["visits"]] == ["Screening", "Baseline", "Visit 3", "Visit 4", "Follow-up"]
    v3 = m["visits"][2]
    assert v3["timing"]["label"] == "Week 12" and v3["window"]["minus"] == 3
    assert m["visits"][4]["window"] is None  # not stated → not invented
    assert any(i["type"] == "window_not_specified" for i in m["issues"])
    hba1c = next(a for a in m["assessments"] if a["name"] == "HbA1c")
    assert hba1c["category"] == "Laboratory" and hba1c["source"]["block_id"]
    assert m["cells"][f"{hba1c['id']}|V4"]["footnotes"] == ["a"]
    assert m["footnotes"][0]["marker"] == "a" and m["conventions"]
    assert body["source_diff"] == []

    # Edit: duplicate timing creates a conflict; revision history is kept and the diff vs source is shown
    m["visits"].append({"id": "V6", "label": "Unplanned week 12 check", "timing_text": "Week 12"})
    del m["cells"][f"{hba1c['id']}|V1"]
    r = client.put(f"/api/design/versions/{v['id']}/soa", json={"matrix": m, "note": "Added check", "base_revision": 1}, headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["revision"] == 2 and len(body["history"]) == 2
    assert {d["type"] for d in body["source_diff"]} >= {"visit_added", "cell_changed"}
    conflicts = [i for i in body["data"]["issues"] if i["severity"] == "conflict"]
    assert {i["type"] for i in conflicts} == {"conflicting_visit", "timing_order"}  # same timing as V3, and out of order after V5
    aid = body["id"]

    # Approval is blocked until conflicts are reviewed; a reviewed issue stays listed with its resolution
    r = client.post(f"/api/design/artifacts/{aid}/status?study_id={s['id']}", json={"status": "approved"}, headers=h)
    assert r.status_code == 409
    for c in conflicts:
        r = client.post(f"/api/design/versions/{v['id']}/soa/issues/{c['id']}/resolve", json={"note": "Intentional safety check"}, headers=h)
        assert r.status_code == 200
        if c is conflicts[0]:
            r2 = client.post(f"/api/design/artifacts/{r.json()['id']}/status?study_id={s['id']}", json={"status": "approved"}, headers=h)
            assert r2.status_code == 409
    assert all(i["resolution"] for i in r.json()["data"]["issues"] if i["severity"] == "conflict")
    aid = r.json()["id"]
    r = client.post(f"/api/design/artifacts/{aid}/status?study_id={s['id']}", json={"status": "approved"}, headers=h)
    assert r.status_code == 200 and r.json()["status"] == "approved"

    # Invalid edits are rejected
    bad = {**m, "cells": {"A99|V1": {"mark": "X"}}}
    assert client.put(f"/api/design/versions/{v['id']}/soa", json={"matrix": bad}, headers=h).status_code == 400

    xl = client.get(f"/api/design/versions/{v['id']}/soa/export?format=xlsx", headers=h)
    wb = load_workbook(io.BytesIO(xl.content))
    sheet = wb["Schedule of Assessments"]
    cells = [c.value for row in sheet.iter_rows() for c in row if c.value]
    assert "Week 12" in cells and "Central laboratory sample collected after an overnight fast." in cells
    doc = Document(io.BytesIO(client.get(f"/api/design/versions/{v['id']}/soa/export?format=docx", headers=h).content))
    assert doc.tables and "Footnotes" in [p.text for p in doc.paragraphs]
    pdf = client.get(f"/api/design/versions/{v['id']}/soa/export?format=pdf", headers=h)
    assert pdf.content.startswith(b"%PDF")


def test_soa_matches_endpoint_timepoints_after_endpoint_analysis(client, study_v1):
    """Endpoints analysed before the SoA exists must still match Week 24 to the Week 24 visit."""
    _, v, h = study_v1
    client.post(f"/api/design/versions/{v['id']}/endpoints/analyze", json={"use_ai": False}, headers=h)
    m = client.post(f"/api/design/versions/{v['id']}/soa/generate", json={"use_ai": False}, headers=h).json()["data"]
    assert not [i for i in m["issues"] if i["type"] == "endpoint_timepoint_unscheduled"]
    # Removing the Week 24 visit makes the endpoint time point unscheduled
    m["visits"] = [x for x in m["visits"] if not (x.get("timing") or {}).get("value") == 24]
    m["cells"] = {k: c for k, c in m["cells"].items() if not k.endswith("|V4")}
    issues = client.put(f"/api/design/versions/{v['id']}/soa", json={"matrix": m}, headers=h).json()["data"]["issues"]
    assert any(i["type"] == "endpoint_timepoint_unscheduled" and "EP1" in i["description"] for i in issues)


# ─── Feature 10: Sample Size & Power ─────────────────────────────────

def test_sample_size_scenarios(client, study_v1):
    s, v, h = study_v1
    inputs = {"endpoint_type": "continuous", "delta": 0.5, "sd": 1, "alpha": 0.05, "power": 0.8, "sided": 2, "dropout": 0.1}
    r1 = client.post(f"/api/design/versions/{v['id']}/sample-size", json={"name": "Base case", "inputs": inputs}, headers=h)
    r2 = client.post(f"/api/design/versions/{v['id']}/sample-size", json={"name": "Higher power", "inputs": {**inputs, "power": 0.9}}, headers=h)
    assert r1.status_code == r2.status_code == 200
    assert r1.json()["data"]["result"]["enrollment"]["total"] == 144
    rows = client.get(f"/api/design/versions/{v['id']}/sample-size", headers=h).json()["scenarios"]
    assert {x["data"]["name"] for x in rows} == {"Base case", "Higher power"}  # neither overwrote the other
    a1 = r1.json()["id"]
    ver = client.post(f"/api/design/versions/{v['id']}/sample-size/{a1}/verify", headers=h).json()
    assert ver["reproducible"]
    r = client.post(f"/api/design/versions/{v['id']}/sample-size", json={"inputs": {**inputs, "design": "crossover"}}, headers=h)
    assert r.status_code == 422 and "Crossover" in r.json()["detail"]
    doc = Document(io.BytesIO(client.get(f"/api/design/versions/{v['id']}/sample-size/{a1}/report?format=docx", headers=h).content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Planning estimate" in text and "N_enroll" in text
    assert client.get(f"/api/design/versions/{v['id']}/sample-size/{a1}/report?format=pdf", headers=h).content.startswith(b"%PDF")


# ─── Feature 11: Protocol Version Control ────────────────────────────

def test_version_registration_and_comparison(client, study_v1):
    s, v1, h = study_v1
    pi_before = client.get(f"/api/design/versions/{v1['id']}/intelligence", headers=h).json()
    assert upload(client, h, s["id"], build_docx("v2"), version_label="1.0").status_code == 400  # label reuse never overwrites
    r = upload(client, h, s["id"], build_docx("v2"), version_label="2.0", amendment_number="A1", effective_date="2026-09-01",
               amendment_reason="Broaden eligibility")
    v2 = r.json()
    assert v2["sequence"] == 2 and v2["amendment_number"] == "A1"
    study = client.get(f"/api/design/studies/{s['id']}", headers=h).json()
    assert [x["label"] for x in study["versions"]] == ["1.0", "2.0"]
    assert study["versions"][0]["file_sha256"] == v1["file_sha256"]
    assert client.get(f"/api/design/versions/{v1['id']}/intelligence", headers=h).json() == pi_before

    r = client.post(f"/api/design/versions/{v2['id']}/compare", json={}, headers=h)
    assert r.status_code == 200, r.text
    d = r.json()["data"]
    assert d["base_version_id"] == v1["id"] and d["match"]["level"] == "confident"
    inc = next(c for c in d["changes"] if c["section_new"] == "4.1 Inclusion Criteria")
    assert inc["type"] == "modified" and "18 to 75" in inc["old"]["text"] and "18 to 80" in inc["new"]["text"]
    assert any(i["area"] == "eligibility" for i in inc["impacts"])
    assert all("recommendation" in i["note"] for i in inc["impacts"])
    assert any(c["category"] == "formatting" and not c["substantive"] for c in d["changes"])
    assert any(c["type"] == "added" and "pancreatitis" in c["new"]["text"] for c in d["changes"])
    areas = {a["area"]: a for a in d["impact_areas"]}
    assert {"eligibility", "sample_size", "schedule"} <= set(areas)

    cid, aid = inc["id"], r.json()["id"]
    eh = add_member(client, h, s["id"], "editor")
    url = f"/api/design/comparisons/{aid}/changes/{cid}/review?study_id={s['id']}"
    assert client.post(url, json={"action": "accept"}, headers=eh).status_code == 403
    assert client.post(url, json={"action": "comment", "comment": "Check recruitment impact"}, headers=eh).status_code == 200
    assert client.post(url, json={"action": "add_action", "action_item": {"description": "Update eligibility checklist", "assignee": "CTM"}},
                       headers=eh).status_code == 200
    r = client.post(url, json={"action": "accept", "comment": "Confirmed"}, headers=h)
    ch = next(c for c in r.json()["data"]["changes"] if c["id"] == cid)
    assert ch["review"]["status"] == "accepted" and ch["comments"] and ch["actions"]
    assert len(r.json()["history"]) == 4


def test_uncertain_document_match_is_flagged(client, study_v1):
    s, v1, h = study_v1
    v3 = upload(client, h, s["id"], build_docx("unrelated"), version_label="3.0").json()
    r = client.post(f"/api/design/versions/{v3['id']}/compare", json={"base_version_id": v1["id"]}, headers=h)
    assert r.status_code == 409
    assert r.json()["detail"]["match"]["level"] in ("uncertain", "unlikely")
    r = client.post(f"/api/design/versions/{v3['id']}/compare", json={"base_version_id": v1["id"], "confirm_match": True}, headers=h)
    assert r.status_code == 200 and r.json()["data"]["match"]["confirmed_by"]


# ─── Feature 12: Reports & Assistant ─────────────────────────────────

def test_reports_use_selected_version(client, study_v1):
    s, v1, h = study_v1
    r = client.post(f"/api/design/versions/{v1['id']}/reports", json={"report_type": "protocol_intelligence"}, headers=h)
    assert r.status_code == 200, r.text
    rep = r.json()["data"]
    assert rep["metadata"]["protocol_version"] == "1.0" and rep["metadata"]["study_id"] == "DX-301"
    assert rep["metadata"]["analysis_versions"][0]["kind"] == "protocol_intelligence"
    items = [it for sec in rep["sections"] for it in sec["items"]]
    phase = next(it for it in items if it["text"].startswith("Study phase"))
    assert phase["label"] == "pending" and phase["citations"][0]
    assert any(it["label"] == "unresolved" and "hypotheses" in it["text"].lower() for it in items)

    v2 = upload(client, h, s["id"], build_docx("v2"), version_label="2.0").json()
    other_pi = client.get(f"/api/design/versions/{v2['id']}/intelligence", headers=h).json()["id"]
    r = client.post(f"/api/design/versions/{v1['id']}/reports", json={"report_type": "protocol_intelligence", "artifact_ids": [other_pi]}, headers=h)
    assert r.status_code == 400 and "different protocol version" in r.json()["detail"]

    client.post(f"/api/design/versions/{v2['id']}/compare", json={}, headers=h)
    r = client.post(f"/api/design/versions/{v2['id']}/reports", json={"report_type": "amendment_impact"}, headers=h)
    rep2 = r.json()
    assert rep2["data"]["metadata"]["protocol_version"] == "2.0"
    assert any("Inclusion" in it["text"] or "18 to 80" in it["text"] for sec in rep2["data"]["sections"] for it in sec["items"])

    # Draft → reviewed → approved with role checks
    aid = rep2["id"]
    eh = add_member(client, h, s["id"], "editor")
    assert client.post(f"/api/design/artifacts/{aid}/status?study_id={s['id']}", json={"status": "approved"}, headers=eh).status_code == 403
    r = client.post(f"/api/design/artifacts/{aid}/status?study_id={s['id']}", json={"status": "reviewed"}, headers=h)
    assert r.json()["status"] == "reviewed"
    doc = Document(io.BytesIO(client.get(f"/api/design/reports/{r.json()['id']}/export?study_id={s['id']}&format=docx", headers=h).content))
    meta = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "DX-301" in meta and "2.0" in meta and "REVIEWED" in meta
    assert client.get(f"/api/design/reports/{aid}/export?study_id={s['id']}&format=pdf", headers=h).content.startswith(b"%PDF")


def test_sample_size_report_matches_saved_calculation(client, study_v1):
    s, v, h = study_v1
    inputs = {"endpoint_type": "binary", "p1": 0.6, "p2": 0.8, "alpha": 0.05, "power": 0.8}
    client.post(f"/api/design/versions/{v['id']}/sample-size", json={"name": "Fleiss", "inputs": inputs}, headers=h)
    rep = client.post(f"/api/design/versions/{v['id']}/reports", json={"report_type": "feasibility_sample_size"}, headers=h).json()["data"]
    calc = next(it for sec in rep["sections"] for it in sec["items"] if it["label"] == "calculated")
    assert "evaluable 164" in calc["text"] and "per arm 82, 82" in calc["text"]


def test_assistant_without_ai_returns_study_evidence_only(client, study_v1):
    _, v, h = study_v1
    r = client.post(f"/api/design/versions/{v['id']}/assistant", json={"question": "What are the inclusion criteria?"}, headers=h)
    body = r.json()
    assert body["outcome"] == "evidence_only" and body["evidence"]
    assert any("18 to 75" in e["text"] for e in body["evidence"])
    assert body["context"]["protocol_version"] == "1.0"
    r = client.post(f"/api/design/versions/{v['id']}/assistant", json={"question": "zzqx quantum blockchain"}, headers=h)
    assert r.json()["outcome"] == "insufficient_evidence" and r.json()["evidence"] == []


def test_assistant_strips_unsupported_claims(client, study_v1, monkeypatch):
    _, v, h = study_v1

    async def fake(system, user, **kw):
        return {"answer": "The study is Phase 3 [E1]. It enrolls 5000 patients [E99].",
                "claims": [{"text": "The study is Phase 3", "citations": ["E1"]},
                           {"text": "It enrolls 5000 patients", "citations": ["E99"]}],
                "limitations": ""}
    monkeypatch.setattr(reports_service, "ai_available", lambda: True)
    monkeypatch.setattr(reports_service, "chat_completion_json", fake)
    body = client.post(f"/api/design/versions/{v['id']}/assistant", json={"question": "What phase is the study?"}, headers=h).json()
    assert body["outcome"] == "answered"
    assert [c["text"] for c in body["claims"]] == ["The study is Phase 3"]
    assert body["claims"][0]["category"] in ("source_fact_pending", "verified_fact")
    assert "[E99]" not in body["answer"] and "removed" in body["limitations"]
    hist = client.get(f"/api/design/versions/{v['id']}/assistant/history", headers=h).json()["items"]
    assert hist[-1]["question"] == "What phase is the study?"
