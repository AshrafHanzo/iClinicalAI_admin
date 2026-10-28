"""Regression tests for the DESIGN scope items completed after review against real protocols."""

import asyncio
import io
import json

from docx import Document
from openpyxl import load_workbook

import services.design_reports as reports_service
from services import endpoint_intelligence as ep
from services import protocol_intelligence as pi
from services import soa_engine
from services.protocol_processing import expand_bullets, process_protocol
from tests.conftest import make_user
from tests.fixtures import (build_complex_pdf, build_docx, build_sap_docx, build_soa_template_docx, build_soa_template_xlsx,
                            build_synopsis_docx_with_hypothesis, build_titled_soa_docx)
from tests.test_design_api import DOCX, add_member, create_study, upload


def extract(blocks, **kw):
    return pi.field_map(asyncio.run(pi.extract_fields(blocks, use_ai=False, **kw)))


# ─── Feature 01: structure and extraction on real-protocol patterns ──

def test_pdf_contents_pages_and_running_headers_are_excluded():
    p = process_protocol("c.pdf", build_complex_pdf())
    assert any(b["kind"] == "toc" for b in p.blocks)
    assert not any("QRS-456,001" in b["text"] for b in p.blocks if b["kind"] != "toc")
    f = extract(p.blocks)
    # Randomization comes from its section, not the contents page line "3.1 Randomization Scheme ....... 3"
    assert f["randomization"]["value"].startswith("Subjects will be randomized in a 4:1")
    assert f["randomization"]["evidence"][0]["page"] == 3


def test_pdf_tables_are_extracted_with_shading():
    p = process_protocol("c.pdf", build_complex_pdf())
    table = next(b for b in p.blocks if b["kind"] == "table")
    assert table["page"] == 5 and table["table"][0][0] == "Assessment"
    assert table["shaded"][1][2] is True and table["shaded"][1][1] is False


def test_symbol_font_bullets_are_split_into_items():
    assert expand_bullets(["The evaluation covers: \x00 Incidence of AEs \x00 Deaths"]) == [
        "The evaluation covers:", "• Incidence of AEs", "• Deaths"]


def test_criteria_under_numbered_subsections_are_collected():
    f = extract(process_protocol("c.pdf", build_complex_pdf()).blocks)
    assert f["inclusion_criteria"]["value"] == ["Age 18 to 55 years.", "Body mass index 18 to 30 kg/m2.", "Stable renal function."]


def test_sample_size_ignores_codes_breakdowns_and_per_cohort_counts():
    f = extract(process_protocol("c.pdf", build_complex_pdf()).blocks)
    # Not 456001 (drug code), 8 (per cohort) or 14/8 (breakdown); "Forty (40)" is the planned total
    assert f["sample_size"]["status"] == "extracted" and f["sample_size"]["value"] == "40"
    assert pi.conflict_key("sample_size", "Approximately 22 subjects (14 HD subjects, 8 healthy subjects) will be enrolled") == "22"


def test_endpoints_only_in_analysis_section_are_flagged():
    f = extract(process_protocol("c.pdf", build_complex_pdf()).blocks)
    pe = f["primary_endpoints"]
    assert pe["value"] == ["Incidence of adverse events.", "Clinical laboratory abnormalities."]
    assert pe["confidence_level"] == "low" and any("analysis section" in x for x in pe["flags"])


def test_background_discussion_is_not_the_study_hypothesis():
    f = extract(process_protocol("c.pdf", build_complex_pdf()).blocks)
    assert f["hypotheses"]["status"] == "missing"


def test_blinding_uses_the_design_term():
    f = extract(process_protocol("c.pdf", build_complex_pdf()).blocks)
    assert f["blinding"]["value"] == "double-blind"


def test_extraction_preferences_limit_fields_and_threshold():
    blocks = process_protocol("p.docx", build_docx()).blocks
    data = asyncio.run(pi.extract_fields(blocks, use_ai=False, preferences={"field_groups": ["Eligibility"], "review_threshold": 0.95}))
    f = pi.field_map(data)
    assert f["inclusion_criteria"]["status"] == "extracted"
    assert f["phase"]["status"] == "not_requested" and f["phase"]["value"].startswith("Not extracted")
    assert data["summary"]["not_requested"] == 23
    assert f["inclusion_criteria"]["needs_review"]  # 0.8 confidence is below the 0.95 threshold
    assert data["preferences"]["review_threshold"] == 0.95


def test_study_characteristics_and_new_fields():
    data = asyncio.run(pi.extract_fields(process_protocol("p.docx", build_docx()).blocks, use_ai=False))
    chars = {c["key"]: c["value"] for c in data["characteristics"]}
    assert chars["phase"] == "Phase 3" and chars["sample_size"] == "240" and chars["study_duration"] == "Not specified"
    f = pi.field_map(data)
    assert "exploratory_objectives" in f and "unclassified_endpoints" in f
    assert f["unclassified_endpoints"]["status"] == "missing"  # endpoints are classified in this protocol


def test_ai_scope_prioritises_relevant_sections_for_long_documents():
    blocks = [{"id": f"x{i}", "kind": "paragraph", "page": i, "section": "9 Appendix", "text": "filler text " * 200} for i in range(90)]
    blocks.append({"id": "ep", "kind": "list_item", "page": 99, "section": "8.2 Primary Endpoints", "text": "Change in FEV1 at Week 12"})
    payload, truncated = pi._blocks_for_ai(blocks, ["primary_endpoints"], "relevant")
    assert "[ep |" in payload  # late endpoint section is included before unrelated filler
    assert truncated is False


# ─── Feature 07: schedule tables from real layouts ───────────────────

def test_shaded_pdf_schedule_is_read_and_needs_confirmation():
    m = asyncio.run(soa_engine.generate(process_protocol("c.pdf", build_complex_pdf()).blocks, use_ai=False))
    assert m["assignment_basis"] == "shading"
    grid = {a["name"]: sorted(k.split("|")[1] for k in m["cells"] if k.startswith(a["id"] + "|")) for a in m["assessments"]}
    assert grid == {"Informed Consent": ["V1"], "Vital Signs": ["V1", "V2", "V3"], "12-lead ECG": ["V2", "V3"], "Urinalysis": ["V3"]}
    vs = next(a for a in m["assessments"] if a["name"] == "Vital Signs")
    assert m["cells"][f"{vs['id']}|V2"]["mark"] == "pre-dose"
    conv = next(i for i in m["issues"] if i["type"] == "shading_convention")
    assert conv["severity"] == "conflict"  # blocks approval until a reviewer confirms
    day8 = m["visits"][2]
    assert day8["window"]["minus"] == 1 and day8["window"]["unit_from_visit_timing"]
    assert any(i["type"] == "window_unit_not_stated" for i in m["issues"])


def test_titled_schedule_with_weekday_row_and_legend():
    for shaded in (False, True):
        m = asyncio.run(soa_engine.generate(process_protocol("t.docx", build_titled_soa_docx(shaded)).blocks, use_ai=False))
        labels = [v["label"] for v in m["visits"]]
        assert labels == ["Screening (Day -28 to Day -2)", "Day -1", "Day 1", "End of Study Visit"]  # Day -1 and Day 1 stay distinct
        names = [a["name"] for a in m["assessments"]]
        assert "Week Day" not in names and "Medical history" in names
        assert next(a for a in m["assessments"] if a["name"] == "Medical history")["footnotes"] == ["b"]
        assert [f["marker"] for f in m["footnotes"]] == ["A", "B"]
        assert m["assignment_basis"] == ("shading" if shaded else "marks")


# ─── Feature 06: hypotheses, SAP, unresolved questions ───────────────

def _pi_with_hypothesis():
    blocks = process_protocol("p.docx", build_docx()).blocks
    data = asyncio.run(pi.extract_fields(blocks, use_ai=False))
    f = pi.field_map(data)
    f["hypotheses"].update({"value": "Drug X is superior to placebo in change from baseline in HbA1c. Drug X improves quality of life questionnaire scores.",
                            "status": "extracted", "evidence": [{"block_id": "b7", "page": None, "section": "2.1", "quote": "x", "verified": True}]})
    return data, blocks


def test_hypotheses_are_linked_to_endpoints():
    data, blocks = _pi_with_hypothesis()
    r = ep.analyze(data, blocks)
    assert [h["id"] for h in r["hypotheses"]] == ["H1", "H2"]
    assert any(l["hypothesis_id"] == "H1" and l["endpoint_id"] == "EP1" for l in r["hypothesis_mapping"])
    assert any(i["type"] == "hypothesis_without_endpoint" and i["hypothesis_id"] == "H2" for i in r["issues"])


def test_sap_comparison_detects_timepoint_mismatch():
    data, blocks = _pi_with_hypothesis()
    sap = process_protocol("sap.docx", build_sap_docx(26)).blocks
    r = ep.analyze(data, blocks, None, sap)
    ep1 = next(x for x in r["sap_comparison"] if x["endpoint_id"] == "EP1")
    assert ep1["status"] == "timepoint_mismatch" and "week 26" in ep1["sap_time_points"]
    assert any(i["type"] == "sap_timepoint_mismatch" and i["certainty"] == "confirmed" for i in r["issues"])
    assert any(x["status"] == "not_found" for x in r["sap_comparison"])  # e.g. body weight is not analysed in the SAP
    assert r["basis"]["sap_used"]


def test_unresolved_questions_are_a_separate_output():
    data, blocks = _pi_with_hypothesis()
    r = ep.analyze(data, blocks)
    qs = r["unresolved_questions"]
    assert qs and all(q["issue_id"] in {i["id"] for i in r["issues"]} for q in qs)
    assert any(q["question"] == "At which time point(s) is EP3 assessed?" for q in qs)


def test_unclassified_endpoints_are_flagged_for_priority():
    data = asyncio.run(pi.extract_fields(process_protocol("p.docx", build_docx()).blocks, use_ai=False))
    f = pi.field_map(data)
    for k in ("primary_endpoints", "secondary_endpoints"):
        f[k].update({"status": "missing", "value": [], "evidence": []})
    f["unclassified_endpoints"].update({"status": "ambiguous", "value": ["Plasma concentration of Drug X at Day 1"],
                                        "evidence": [{"block_id": "b12", "item": 0, "verified": True}]})
    r = ep.analyze(data, [])
    assert r["register"][0]["type"] == "unclassified"
    assert any(i["type"] == "priority_not_stated" for i in r["issues"])


# ─── Feature 12: assistant contradiction handling ────────────────────

def test_assistant_reports_contradictions_with_source_to_review():
    blocks = process_protocol("p.docx", build_docx(synopsis_sample_size="Approximately 200 participants")).blocks
    data = asyncio.run(pi.extract_fields(blocks, use_ai=False))
    corpus = reports_service.build_corpus(blocks, {"protocol_intelligence": [{"data": data}]}, {}, False)
    res = asyncio.run(reports_service.answer("What is the planned sample size?", corpus))
    contra = [i for i in res["review_items"] if i["type"] == "contradictory"]
    assert contra and "Planned sample size" in contra[0]["source_to_review"]
    assert "Contradictory evidence" in res["limitations"]


def test_assistant_names_missing_source():
    blocks = process_protocol("p.docx", build_docx()).blocks
    data = asyncio.run(pi.extract_fields(blocks, use_ai=False))
    corpus = reports_service.build_corpus(blocks, {"protocol_intelligence": [{"data": data}]}, {}, False)
    res = asyncio.run(reports_service.answer("What is the study duration?", corpus))
    assert any(i["type"] == "missing" and "Study duration" in i["source_to_review"] for i in res["review_items"])


# ─── API: preferences, supporting documents, templates, reports ──────

def test_supporting_document_feeds_extraction(client, owner):
    _, h = owner
    s = create_study(client, h, "SUP-1")
    files = [("file", ("p.docx", build_docx(), DOCX)), ("supporting_files", ("syn.docx", build_synopsis_docx_with_hypothesis(), DOCX))]
    r = client.post(f"/api/design/studies/{s['id']}/versions", files=files,
                    data={"preferences": json.dumps({"use_ai": False}), "supporting_kinds": "synopsis", "version_label": "1"}, headers=h)
    assert r.status_code == 200, r.text
    vid = r.json()["id"]
    hyp = {f["key"]: f for f in client.get(f"/api/design/versions/{vid}/intelligence", headers=h).json()["data"]["fields"]}["hypotheses"]
    assert hyp["status"] == "extracted" and hyp["evidence"][0]["block_id"].startswith("s1-")
    # The cited supporting-document passage can be opened in the evidence viewer
    blk = client.get(f"/api/design/versions/{vid}/blocks/{hyp['evidence'][0]['block_id']}", headers=h)
    assert blk.status_code == 200 and "superior to placebo" in blk.json()["block"]["text"]

    r = client.post(f"/api/design/studies/{s['id']}/versions", files=files,
                    data={"preferences": json.dumps({"use_ai": False, "include_supporting": False}),
                          "supporting_kinds": "synopsis", "version_label": "2"}, headers=h)
    hyp2 = {f["key"]: f for f in client.get(f"/api/design/versions/{r.json()['id']}/intelligence", headers=h).json()["data"]["fields"]}["hypotheses"]
    assert hyp2["status"] == "missing"


def test_preferences_are_applied_on_upload(client, owner):
    _, h = owner
    s = create_study(client, h, "PREF-1")
    r = upload(client, h, s["id"], build_docx(), preferences=json.dumps({"use_ai": False, "field_groups": ["Endpoints", "Objectives"]}))
    data = client.get(f"/api/design/versions/{r.json()['id']}/intelligence", headers=h).json()["data"]
    f = {x["key"]: x for x in data["fields"]}
    assert f["primary_endpoints"]["status"] == "extracted" and f["inclusion_criteria"]["status"] == "not_requested"
    assert data["preferences"]["field_groups"] == ["Endpoints", "Objectives"]
    wb = load_workbook(io.BytesIO(client.get(f"/api/design/versions/{r.json()['id']}/intelligence/export?format=xlsx", headers=h).content))
    assert "Study characteristics" in wb.sheetnames


def test_endpoint_analysis_uses_uploaded_sap(client, owner):
    _, h = owner
    s = create_study(client, h, "SAP-1")
    files = [("file", ("p.docx", build_docx(), DOCX)), ("supporting_files", ("sap.docx", build_sap_docx(26), DOCX))]
    r = client.post(f"/api/design/studies/{s['id']}/versions", files=files,
                    data={"preferences": json.dumps({"use_ai": False, "include_supporting": False}),
                          "supporting_kinds": "statistical_analysis_plan"}, headers=h)
    d = client.post(f"/api/design/versions/{r.json()['id']}/endpoints/analyze", json={"use_ai": False}, headers=h).json()["data"]
    assert d["basis"]["sap_used"] and any(i["type"] == "sap_timepoint_mismatch" for i in d["issues"])
    assert d["unresolved_questions"]


def test_soa_templates_and_timeline(client, owner):
    _, h = owner
    s = create_study(client, h, "TPL-1")
    v = upload(client, h, s["id"], build_docx(), version_label="1.0").json()
    soa = client.post(f"/api/design/versions/{v['id']}/soa/generate", json={"use_ai": False}, headers=h).json()
    assert [t["label"].split(" / ")[0] for t in soa["timeline"]][0] == "Screening"
    assert soa["timeline"][3]["approx_day"] == 168  # Week 24, derived only for ordering

    bad = client.post(f"/api/design/studies/{s['id']}/soa-templates", files={"file": ("t.txt", b"x", "text/plain")}, headers=h)
    assert bad.status_code == 400
    tx = client.post(f"/api/design/studies/{s['id']}/soa-templates", files={"file": ("sponsor.xlsx", build_soa_template_xlsx(), "application/octet-stream")},
                     data={"name": "Sponsor Excel"}, headers=h).json()
    assert "{{SOA_TABLE}}" in tx["data"]["placeholders"] and "path" not in tx["data"]
    td = client.post(f"/api/design/studies/{s['id']}/soa-templates", files={"file": ("sponsor.docx", build_soa_template_docx(), DOCX)}, headers=h).json()
    assert len(client.get(f"/api/design/studies/{s['id']}/soa-templates", headers=h).json()["templates"]) == 2

    r = client.get(f"/api/design/versions/{v['id']}/soa/export?format=xlsx&template_id={tx['id']}", headers=h)
    ws = load_workbook(io.BytesIO(r.content))["Sponsor SoA"]
    assert ws["A1"].value == "Sponsor SoA template — Study TPL-1 version 1.0 (DRAFT)"
    assert ws["A3"].value == "Assessment" and ws["B3"].value.startswith("Screening")
    assert ws["H3"].value.startswith("a:")
    doc = Document(io.BytesIO(client.get(f"/api/design/versions/{v['id']}/soa/export?format=docx&template_id={td['id']}", headers=h).content))
    assert doc.paragraphs[0].text == "Sponsor SoA — TPL-1 / 1.0" and doc.tables and doc.tables[0].rows[0].cells[0].text == "Assessment"
    assert client.get(f"/api/design/versions/{v['id']}/soa/export?format=docx&template_id={tx['id']}", headers=h).status_code == 400
    wb = load_workbook(io.BytesIO(client.get(f"/api/design/versions/{v['id']}/soa/export?format=xlsx", headers=h).content))
    assert "Timeline" in wb.sheetnames


def test_custom_report_templates_and_explicit_sources(client, owner):
    _, h = owner
    s = create_study(client, h, "REP-1")
    v1 = upload(client, h, s["id"], build_docx(), version_label="1.0").json()
    v2 = upload(client, h, s["id"], build_docx("v2"), version_label="2.0").json()
    client.post(f"/api/design/versions/{v2['id']}/compare", json={}, headers=h)

    cat = client.get("/api/design/reports/types", headers=h).json()["sections"]
    assert "synopsis" in cat and "version_history" in cat and "endpoint_questions" in cat
    eh = add_member(client, h, s["id"], "editor")
    vh = add_member(client, h, s["id"], "viewer")
    body = {"name": "Board pack", "title": "Amendment board pack", "intro": "Prepared for the study board.",
            "sections": ["version_history", "characteristics", "synopsis", "impacts"]}
    assert client.post(f"/api/design/studies/{s['id']}/report-templates", json=body, headers=vh).status_code == 403
    assert client.post(f"/api/design/studies/{s['id']}/report-templates", json={**body, "sections": ["nope"]}, headers=eh).status_code == 400
    t = client.post(f"/api/design/studies/{s['id']}/report-templates", json=body, headers=eh).json()
    t2 = client.put(f"/api/design/studies/{s['id']}/report-templates/{t['id']}", json={**body, "title": "Amendment board pack v2"}, headers=eh).json()
    assert t2["revision"] == 2 and t2["lineage_id"] == t["lineage_id"]

    rep = client.post(f"/api/design/versions/{v2['id']}/reports", json={"template_id": t2["id"]}, headers=h).json()["data"]
    assert rep["title"] == "Amendment board pack v2" and rep["report_type"] == "custom" and rep["template"]["revision"] == 2
    headings = [x["heading"] for x in rep["sections"]]
    assert headings[:3] == ["Introduction", "Protocol version history", "Study characteristics"]
    hist = next(x for x in rep["sections"] if x["heading"] == "Protocol version history")
    assert [r[1] for r in hist["table"]["rows"]] == ["1.0", "2.0"]
    syn = next(x for x in rep["sections"] if x["heading"].startswith("Study synopsis"))
    assert "Synopsis Generator has not been run" in syn["items"][0]["text"]

    src = client.get(f"/api/design/versions/{v1['id']}/report-sources", headers=h).json()
    pi_id = next(a["id"] for a in src["artifacts"] if a["kind"] == "protocol_intelligence")
    assert any(e["key"] == "summary" for e in src["existing_analyses"])
    r = client.post(f"/api/design/versions/{v1['id']}/reports", json={"report_type": "protocol_intelligence", "artifact_ids": [pi_id]}, headers=h).json()
    assert [x["artifact_id"] for x in r["data"]["metadata"]["analysis_versions"] if "artifact_id" in x] == [pi_id]
    r = client.post(f"/api/design/versions/{v1['id']}/reports", json={"report_type": "protocol_intelligence", "artifact_ids": []}, headers=h).json()
    assert "Protocol Intelligence has not been run" in json.dumps(r["data"]["sections"])  # nothing chosen, nothing used
    doc = Document(io.BytesIO(client.get(f"/api/design/reports/{r['id']}/export?study_id={s['id']}&format=docx", headers=h).content))
    assert doc.paragraphs


def test_other_users_cannot_use_study_templates(client, owner):
    _, h = owner
    s = create_study(client, h, "ISO-2")
    _, oh = make_user()
    assert client.get(f"/api/design/studies/{s['id']}/report-templates", headers=oh).status_code == 404
    assert client.get(f"/api/design/studies/{s['id']}/soa-templates", headers=oh).status_code == 404
