"""Protocol Review: completeness, consistency, ambiguity, unclear wording and unresolved issues,
with severity, priority, recommendations, source references and the resolve/document workflow."""

import asyncio
import io

from docx import Document
from openpyxl import load_workbook

import services.protocol_review as review_service
from services import protocol_intelligence as pi
from services.protocol_processing import process_protocol
from tests.fixtures import build_docx
from tests.test_design_api import add_member, create_study, study_v1, upload  # noqa: F401  (fixture)


def _review(content: bytes, name: str = "p.docx") -> dict:
    blocks = process_protocol(name, content).blocks
    data = asyncio.run(pi.extract_fields(blocks, use_ai=False))
    return asyncio.run(review_service.review(data, blocks, use_ai=False))


def _docx(paragraphs: list[tuple[str, str]]) -> bytes:
    d = Document()
    for kind, text in paragraphs:
        if kind == "h1":
            d.add_heading(text, level=1)
        elif kind == "h2":
            d.add_heading(text, level=2)
        elif kind == "li":
            d.add_paragraph(text, style="List Number")
        else:
            d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_every_finding_has_the_required_output_fields():
    r = _review(build_docx())
    assert r["findings"]
    for f in r["findings"]:
        assert f["category"] in review_service.CATEGORIES and f["category_label"]
        assert f["severity"] in ("critical", "major", "minor")
        assert f["priority"] == {"critical": "high", "major": "medium", "minor": "low"}[f["severity"]]
        assert f["title"] and f["description"] and f["recommendation"]
        assert f["status"] == "open" and f["origin"] in ("rule", "ai")
        if f["category"] != "completeness":
            assert f["evidence"], f  # issues found in the text cite where they are


def test_inconsistent_primary_endpoint_timepoint_is_critical():
    # Fixture: primary endpoint at Week 24 in the endpoint section, Week 26 in the statistics section
    r = _review(build_docx())
    f = next(f for f in r["findings"] if f["category"] == "consistency" and "time point" in f["title"])
    assert f["severity"] == "critical"
    assert {e["section"] for e in f["evidence"]} >= {"2.3 Primary Endpoint", "7 Statistical Methods"}


def test_missing_information_and_sections_are_reported():
    content = _docx([("h1", "1 Study Design"), ("p", "This is an open-label single-arm study of Drug Z."),
                     ("h1", "2 Inclusion Criteria"), ("li", "Adults aged 18 to 65 years."),
                     ("h1", "3 Treatment"), ("p", "Drug Z 5 mg once daily for 12 weeks. " * 3)])
    r = _review(content)
    titles = {f["title"] for f in r["findings"]}
    assert "Primary endpoint not specified" in titles
    assert "Exclusion criteria not specified" in titles
    assert "No section on adverse event definitions and reporting" in titles
    missing = next(f for f in r["findings"] if f["title"] == "Primary endpoint not specified")
    assert missing["severity"] == "critical" and missing["category"] == "completeness"
    comp = r["summary"]["completeness"]
    assert comp["checked"] > comp["present"] and 0 <= comp["percent"] < 100
    assert any(c["element"] == "Primary endpoint" and c["status"] == "missing" for c in r["checklist"])


def test_ambiguous_and_unclear_wording_and_placeholders():
    content = _docx([("h1", "1 Study Design"), ("p", "This is a randomized, double-blind study. Approximately 60 subjects will be enrolled."),
                     ("h1", "2 Inclusion Criteria"),
                     ("li", "Adequate renal function."),
                     ("li", "Subjects should have stable disease."),
                     ("h1", "3 Exclusion Criteria"), ("li", "Any clinically significant ECG abnormality."),
                     ("h1", "4 Primary Endpoint"), ("li", "Change from baseline in FEV1 at Week 12 measured by spirometry."),
                     ("h1", "5 Statistical Methods"), ("p", "The primary analysis model is TBD.")])
    r = _review(content)
    by_cat = {}
    for f in r["findings"]:
        by_cat.setdefault(f["category"], []).append(f)
    ambiguous_text = " ".join(f["description"] for f in by_cat["ambiguity"])
    assert "Adequate" in ambiguous_text and "clinically significant" in ambiguous_text
    assert any("should" in f["description"] for f in by_cat["clarity"])
    placeholder = next(f for f in by_cat["unresolved"] if f["title"] == "Placeholder or undecided content")
    assert placeholder["evidence"][0]["quote"] == "TBD" and placeholder["evidence"][0]["section"] == "5 Statistical Methods"
    # Randomized design without a randomization description is an inconsistency
    assert any(f["key"] == "consistency:randomization_missing" for f in by_cat["consistency"])


def test_non_randomized_title_with_pdf_spacing_is_not_a_contradiction():
    content = _docx([("h1", "A Phase 1, Open -Label, Non -Randomized Study"), ("p", "Subjects will not be randomized."),
                     ("h1", "1 Study Design"), ("p", "This is a non-randomized, open-label study in healthy adults. " * 3)])
    r = _review(content)
    assert not any(f["key"] == "consistency:design_randomization" for f in r["findings"])


def test_standard_deviation_is_not_a_protocol_deviation_section():
    content = _docx([("h1", "1 Statistics"), ("p", "Results are summarised by mean and standard deviation. " * 6)])
    r = _review(content)
    item = next(c for c in r["checklist"] if c["element"] == "Protocol deviations")
    assert item["status"] == "missing"


def test_ai_findings_require_quotes_found_in_the_protocol(monkeypatch):
    blocks = process_protocol("p.docx", build_docx()).blocks
    data = asyncio.run(pi.extract_fields(blocks, use_ai=False))

    async def fake_ai(system, user, **kw):
        return {"findings": [
            {"category": "clarity", "severity": "major", "title": "Dose wording unclear", "description": "The dose text is unclear.",
             "recommendation": "State the dose.", "evidence": [{"block_id": "b21", "quote": "Drug X 10 mg once daily"}]},
            {"category": "ambiguity", "severity": "minor", "title": "Invented issue", "description": "Not in the text.",
             "recommendation": "n/a", "evidence": [{"block_id": "b1", "quote": "this sentence does not exist anywhere"}]},
        ]}
    monkeypatch.setattr(review_service, "ai_available", lambda: True)
    monkeypatch.setattr(review_service, "chat_completion_json", fake_ai)
    r = asyncio.run(review_service.review(data, blocks, use_ai=True))
    ai = [f for f in r["findings"] if f["origin"] == "ai"]
    assert [f["title"] for f in ai] == ["Dose wording unclear"]
    assert ai[0]["evidence"][0]["section"] == "3.1 Treatment Arms" and ai[0]["certainty"] == "potential"
    assert r["basis"]["ai_used"] is True


# ─── API: run, resolve / document workflow, approval, exports ────────

def test_review_api_workflow(client, study_v1):  # noqa: F811
    s, v, h = study_v1
    vid = v["id"]
    assert client.get(f"/api/design/versions/{vid}/review", headers=h).status_code == 404
    r = client.post(f"/api/design/versions/{vid}/review/run", json={"use_ai": False}, headers=h)
    assert r.status_code == 200, r.text
    art = r.json()
    assert art["kind"] == "protocol_review" and art["data"]["basis"]["ai_used"] is False
    critical = [f for f in art["data"]["findings"] if f["severity"] == "critical"]
    assert critical

    # Approval is blocked while critical findings are open
    r = client.post(f"/api/design/artifacts/{art['id']}/status?study_id={s['id']}", json={"status": "approved"}, headers=h)
    assert r.status_code == 409

    # A note is required to resolve or document
    fid = critical[0]["id"]
    assert client.post(f"/api/design/versions/{vid}/review/findings/{fid}", json={"action": "resolve"}, headers=h).status_code == 400
    r = client.post(f"/api/design/versions/{vid}/review/findings/{fid}", json={"action": "resolve", "note": "Fixed in v1.1"}, headers=h)
    assert r.status_code == 200 and r.json()["revision"] == 2
    others = [f for f in r.json()["data"]["findings"] if f["severity"] == "critical" and f["status"] == "open"]
    for f in others:
        r = client.post(f"/api/design/versions/{vid}/review/findings/{f['id']}",
                        json={"action": "document", "note": "Accepted: covered in the SAP"}, headers=h)
        assert r.status_code == 200
    data = r.json()["data"]
    assert next(f for f in data["findings"] if f["id"] == fid)["status"] == "resolved"
    assert data["summary"]["open_critical"] == 0
    latest = r.json()["id"]
    r = client.post(f"/api/design/artifacts/{latest}/status?study_id={s['id']}", json={"status": "approved"}, headers=h)
    assert r.status_code == 200, r.text

    # Re-running keeps review decisions for findings found again
    r = client.post(f"/api/design/versions/{vid}/review/run", json={"use_ai": False}, headers=h)
    again = next(f for f in r.json()["data"]["findings"] if f["key"] == critical[0]["key"])
    assert again["status"] == "resolved" and again["resolution"]["note"] == "Fixed in v1.1"

    # Exports
    x = client.get(f"/api/design/versions/{vid}/review/export?format=xlsx", headers=h)
    wb = load_workbook(io.BytesIO(x.content))
    assert {"Findings", "Completeness checklist"} <= set(wb.sheetnames)
    assert wb["Findings"]["B1"].value == "Severity"
    d = client.get(f"/api/design/versions/{vid}/review/export?format=docx", headers=h)
    text = "\n".join(p.text for p in Document(io.BytesIO(d.content)).paragraphs)
    assert "Recommendation:" in text and "Completeness checklist" in text
    assert client.get(f"/api/design/versions/{vid}/review/export?format=pdf", headers=h).content[:4] == b"%PDF"

    # Reports include the findings
    r = client.post(f"/api/design/versions/{vid}/reports", json={"report_type": "study_design_assessment"}, headers=h)
    headings = [sec["heading"] for sec in r.json()["data"]["sections"]]
    assert "Protocol Review findings" in headings


def test_review_permissions(client, study_v1):  # noqa: F811
    s, v, h = study_v1
    vid = v["id"]
    viewer = add_member(client, h, s["id"], "viewer")
    editor = add_member(client, h, s["id"], "editor")
    assert client.post(f"/api/design/versions/{vid}/review/run", json={"use_ai": False}, headers=viewer).status_code == 403
    r = client.post(f"/api/design/versions/{vid}/review/run", json={"use_ai": False}, headers=editor)
    assert r.status_code == 200
    fid = r.json()["data"]["findings"][0]["id"]
    assert client.get(f"/api/design/versions/{vid}/review", headers=viewer).status_code == 200
    # Editors can comment; resolving needs a reviewer
    assert client.post(f"/api/design/versions/{vid}/review/findings/{fid}", json={"action": "comment", "note": "Check this"},
                       headers=editor).status_code == 200
    assert client.post(f"/api/design/versions/{vid}/review/findings/{fid}", json={"action": "resolve", "note": "x"},
                       headers=editor).status_code == 403
